import re
import os
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN

BASE = os.path.dirname(os.path.abspath(__file__))
SRC = os.path.join(BASE, "ppt_tutores.md")
OUT = os.path.join(BASE, "ppt_tutores.pptx")

BULU = RGBColor(0x2C, 0x6F, 0xBB)
DARK = RGBColor(0x22, 0x2A, 0x35)
ACCENT = RGBColor(0x7A, 0x3F, 0xA0)
GRAY = RGBColor(0x5A, 0x64, 0x72)

prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)
BLANK = prs.slide_layouts[6]

def add_background(slide, color=DARK):
    shp = slide.shapes.add_shape(1, 0, 0, prs.slide_width, prs.slide_height)
    shp.fill.solid()
    shp.fill.fore_color.rgb = color
    shp.line.fill.background()
    shp.shadow.inherit = False
    slide.shapes._spTree.remove(shp._element)
    slide.shapes._spTree.insert(2, shp._element)
    return shp

def add_text(slide, top, left, width, height, text, size=18, color=None, bold=False, align=PP_ALIGN.LEFT):
    tb = slide.shapes.add_textbox(Inches(left), Inches(top), Inches(width), Inches(height))
    tf = tb.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = text
    p.font.size = Pt(size)
    p.font.color.rgb = color if color else RGBColor(0xFF, 0xFF, 0xFF)
    p.font.bold = bold
    p.alignment = align
    return tb

SUMMARY_RE = re.compile(r"^\*\*Resumen:?\*\*\s*(.*)$")

def parse_slides(content):
    slides = []
    cur = None
    for raw in content.splitlines():
        s = raw.strip()
        if not s:
            continue
        m = re.match(r"^## Slide (\d+) — (.+)$", s)
        if m:
            if cur:
                slides.append(cur)
            cur = {"title": m.group(2), "resumen": None, "bullets": []}
            continue
        if not cur:
            continue
        m2 = SUMMARY_RE.match(s)
        if m2:
            cur["resumen"] = m2.group(1).strip()
            continue
        if s.startswith("- "):
            cur["bullets"].append(s[2:].strip())
    if cur:
        slides.append(cur)
    return slides

with open(SRC, encoding="utf-8") as f:
    content = f.read()

slides = parse_slides(content)

def add_title_slide(title):
    slide = prs.slides.add_slide(BLANK)
    add_background(slide, DARK)
    bar = slide.shapes.add_shape(1, Inches(0.9), Inches(2.7), Inches(0.18), Inches(1.6))
    bar.fill.solid()
    bar.fill.fore_color.rgb = BULU
    bar.line.fill.background()
    bar.shadow.inherit = False
    add_text(slide, 2.5, 1.3, 11.0, 1.6, "Detección de anomalías en contrataciones públicas (DNCP Paraguay)", size=34, bold=True, color=RGBColor(0xFF, 0xFF, 0xFF))
    add_text(slide, 3.3, 1.5, 10.5, 2.0, "Del portal OCDS + PDFs de pliegos hasta features por licitación\nPipeline ETL de 7 etapas · datos 2021–2025/2026", size=18, color=GRAY)

def add_content_slide(title, resumen, bullets):
    slide = prs.slides.add_slide(BLANK)
    add_background(slide, RGBColor(0xFF, 0xFF, 0xFF))
    top = slide.shapes.add_shape(1, 0, 0, prs.slide_width, Inches(1.15))
    top.fill.solid()
    top.fill.fore_color.rgb = DARK
    top.line.fill.background()
    top.shadow.inherit = False
    tb = top.text_frame
    tb.margin_left = Inches(0.5)
    tb.margin_top = Inches(0.22)
    p = tb.paragraphs[0]
    p.text = title
    p.font.size = Pt(26)
    p.font.bold = True
    p.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)

    img = None
    full = False
    for b in bullets:
        if b.startswith("GráficaFull:"):
            img = b.replace("GráficaFull:", "").strip().strip("`")
            full = True
            break
        if b.startswith("Gráfica:"):
            img = b.replace("Gráfica:", "").strip().strip("`")
            break
    bullets = [b for b in bullets if not b.startswith("Gráfica")]

    # Placement de imagen
    if img:
        rel = os.path.join(BASE, img)
        if os.path.exists(rel):
            from PIL import Image
            im = Image.open(rel)
            ar = im.width / im.height
            if full:
                max_w, max_h = Inches(12.0), Inches(4.5)
                w = min(max_w, max_h * ar)
                h = w / ar
                left = (prs.slide_width - w) / 2
                top_pic = Inches(2.15) if resumen else Inches(1.35)
                slide.shapes.add_picture(rel, left, top_pic, width=w, height=h)
            else:
                slide.shapes.add_picture(rel, Inches(3.6), Inches(1.6), height=Inches(5.2))
        else:
            add_text(slide, 1.0, 1.5, 11.3, 0.5, f"[Imagen no encontrada: {img}]", size=12, color=RGBColor(0xC0, 0x39, 0x2B))

    text_w = 12.0 if full else (2.7 if img else 12.0)
    if resumen:
        top_res = 1.3 if not full else 1.3
        tb2 = slide.shapes.add_textbox(Inches(0.7), Inches(top_res), Inches(text_w), Inches(0.9)).text_frame
        tb2.word_wrap = True
        p = tb2.paragraphs[0]
        p.text = resumen
        p.font.size = Pt(15)
        p.font.italic = True
        p.font.color.rgb = ACCENT
        p.space_after = Pt(6)

    top_b = 2.45 if resumen else 1.5
    if full:
        top_b = 6.75 if resumen else 6.1
    tb3 = slide.shapes.add_textbox(Inches(0.7), Inches(top_b), Inches(text_w), Inches(7.3 - top_b)).text_frame
    tb3.word_wrap = True
    first = True
    for b in bullets:
        p = tb3.paragraphs[0] if first else tb3.add_paragraph()
        first = False
        p.text = "• " + b
        p.font.size = Pt(15)
        p.font.color.rgb = RGBColor(0x26, 0x2B, 0x33)
        p.space_after = Pt(7)

if slides:
    add_title_slide("Título")
    for s in slides[1:]:
        add_content_slide(s["title"], s.get("resumen"), s["bullets"])

prs.save(OUT)
print(f"Generado: {OUT} — {len(slides)} slides")